from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


MAX_DOCUMENT_BYTES = 8_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 16_000_000
MAX_DOCX_ARCHIVE_ENTRIES = 2_048
MAX_PDF_PAGES = 500
MAX_TEXT_CHARS = 200_000
MAX_TABLES = 100
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 100
MAX_FILENAME_LENGTH = 255
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_FILENAME_CONTROL_CHARS = re.compile(r"[\x00-\x1f\x7f]")


@dataclass(frozen=True)
class DocumentExtraction:
    filename: str
    mime_type: str
    page_count: int | None
    paragraphs: int
    characters: int
    text: str
    truncated: bool
    tables: list[list[list[str]]]


def _limit_text(text: str) -> tuple[str, bool]:
    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(normalized) <= MAX_TEXT_CHARS:
        return normalized, False
    return normalized[:MAX_TEXT_CHARS], True


def _normalize_table(rows: list[list[str]]) -> list[list[str]]:
    width = min(max((len(row) for row in rows), default=0), MAX_TABLE_COLS)
    normalized: list[list[str]] = []
    for row in rows[:MAX_TABLE_ROWS]:
        cells = [" ".join(str(cell).split()) for cell in row[:width]]
        normalized.append(cells + [""] * (width - len(cells)))
    return normalized


def _validate(content: bytes) -> None:
    if not content:
        raise ValueError("document content cannot be empty")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("document exceeds 8 MB limit")


def _safe_filename(filename: str) -> str:
    candidate = Path(filename).name.strip()
    if (
        not candidate
        or candidate in {".", ".."}
        or len(candidate) > MAX_FILENAME_LENGTH
        or _FILENAME_CONTROL_CHARS.search(candidate)
    ):
        raise ValueError("invalid document filename")
    return candidate


def _validate_docx_archive(archive: ZipFile) -> None:
    infos = archive.infolist()
    if len(infos) > MAX_DOCX_ARCHIVE_ENTRIES:
        raise ValueError("invalid DOCX document")
    total_uncompressed = 0
    for info in infos:
        if info.flag_bits & 0x1:
            raise ValueError("invalid DOCX document")
        filename = info.filename.replace("\\", "/")
        parts = [part for part in filename.split("/") if part]
        if filename.startswith("/") or ".." in parts:
            raise ValueError("invalid DOCX document")
        total_uncompressed += max(0, int(info.file_size))
        if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ValueError("DOCX archive expands beyond safety limit")


def _validate_signature(mime_type: str, content: bytes) -> None:
    if mime_type == PDF_MIME:
        if not content.startswith(b"%PDF-"):
            raise ValueError("invalid PDF document")
        return
    if mime_type == DOCX_MIME:
        if not content.startswith(b"PK\x03\x04"):
            raise ValueError("invalid DOCX document")
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                _validate_docx_archive(archive)
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise ValueError("invalid DOCX document")
        except (BadZipFile, OSError) as exc:
            raise ValueError("invalid DOCX document") from exc
        return
    raise ValueError("unsupported document type")


def _extract_pdf(filename: str, content: bytes) -> DocumentExtraction:
    try:
        reader = PdfReader(io.BytesIO(content))
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise ValueError("PDF exceeds page limit")
        pages = [page.extract_text() or "" for page in reader.pages]
    except ValueError:
        raise
    except (PdfReadError, OSError) as exc:
        raise ValueError("invalid PDF document") from exc
    tables: list[list[list[str]]] = []
    for page_text in pages:
        candidate_rows: list[list[str]] = []
        for line in page_text.splitlines():
            raw = line.strip()
            if raw.count("|") >= 1:
                cells = [cell.strip() for cell in raw.split("|")]
            elif "\t" in raw:
                cells = [cell.strip() for cell in raw.split("\t")]
            else:
                continue
            if len(cells) >= 2 and any(cells):
                candidate_rows.append(cells)
        if len(candidate_rows) >= 2 and len(tables) < MAX_TABLES:
            tables.append(_normalize_table(candidate_rows))

    raw_text = "\n\n".join(pages)
    text, truncated = _limit_text(raw_text)
    return DocumentExtraction(
        filename=filename,
        mime_type=PDF_MIME,
        page_count=page_count,
        paragraphs=len([p for p in text.split("\n\n") if p.strip()]),
        characters=len(text),
        text=text,
        truncated=truncated,
        tables=tables,
    )


def extract_document(filename: str, mime_type: str, content: bytes) -> DocumentExtraction:
    """Extract text and structured tables from supported documents."""

    _validate(content)
    filename = _safe_filename(str(filename))
    mime_type = mime_type.strip().lower()
    if mime_type not in {PDF_MIME, DOCX_MIME}:
        raise ValueError("unsupported document type")
    _validate_signature(mime_type, content)

    if mime_type == PDF_MIME:
        return _extract_pdf(filename, content)

    try:
        document = Document(io.BytesIO(content))
    except (BadZipFile, ValueError, OSError) as exc:
        raise ValueError("invalid DOCX document") from exc
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    raw_text = "\n\n".join(paragraphs)
    text, truncated = _limit_text(raw_text)
    tables: list[list[list[str]]] = []
    for table in document.tables[:MAX_TABLES]:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if rows:
            tables.append(_normalize_table(rows))

    return DocumentExtraction(
        filename=filename,
        mime_type=mime_type,
        page_count=None,
        paragraphs=len(paragraphs),
        characters=len(text),
        text=text,
        truncated=truncated,
        tables=tables,
    )
