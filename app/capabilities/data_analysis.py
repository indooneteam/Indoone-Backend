from __future__ import annotations

import csv
import io
import json
import math
from statistics import mean, median, pstdev
from typing import Any

from openpyxl import load_workbook

MAX_ANALYSIS_BYTES = 8_000_000
MAX_SAMPLE_ROWS = 10
MAX_ROWS = 50_000
MAX_COLUMNS = 500


def _numbers(values: list[Any]) -> list[float]:
    numbers: list[float] = []
    for value in values:
        try:
            number = float(value)
        except (TypeError, ValueError):
            continue
        if math.isfinite(number):
            numbers.append(number)
    return numbers


def _summarize_records(records: list[dict[str, Any]]) -> dict[str, Any]:
    columns = sorted({str(key) for row in records for key in row})
    summary: dict[str, Any] = {}
    for column in columns:
        values = [row.get(column) for row in records]
        numbers = _numbers(values)
        non_empty = [value for value in values if value not in (None, "")]
        item: dict[str, Any] = {
            "non_empty": len(non_empty),
            "missing": len(values) - len(non_empty),
        }
        if numbers:
            item.update({
                "numeric": True,
                "count": len(numbers),
                "min": min(numbers),
                "max": max(numbers),
                "mean": mean(numbers),
                "median": median(numbers),
                "stddev": pstdev(numbers) if len(numbers) > 1 else 0.0,
            })
        else:
            item["numeric"] = False
            item["unique_values"] = len({str(value) for value in non_empty})
        summary[column] = item
    return summary


def _duplicate_rows(records: list[dict[str, Any]]) -> int:
    seen: set[str] = set()
    duplicates = 0
    for row in records:
        key = json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)
        if key in seen:
            duplicates += 1
        else:
            seen.add(key)
    return duplicates


def _chart_specs(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    specs: list[dict[str, Any]] = []
    for column, item in _summarize_records(records).items():
        if item.get("numeric"):
            specs.append({
                "type": "histogram",
                "column": column,
                "title": f"Distribution of {column}",
            })
        elif item.get("unique_values", 0) <= 20:
            specs.append({
                "type": "bar",
                "column": column,
                "title": f"Values of {column}",
            })
        if len(specs) >= 10:
            break
    return specs


def _read_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            continue
        normalized_headers = [str(value).strip() if value not in (None, "") else f"column_{index + 1}" for index, value in enumerate(headers)]
        for values in rows:
            if not any(value not in (None, "") for value in values):
                continue
            row = {normalized_headers[index]: values[index] if index < len(values) else None for index in range(len(normalized_headers))}
            row["_sheet"] = worksheet.title
            records.append(row)
            if len(records) >= MAX_ROWS:
                break
        if len(records) >= MAX_ROWS:
            break
    workbook.close()
    return records


def _parse_records(filename: str, content: bytes) -> list[dict[str, Any]]:
    name = filename.casefold()
    if name.endswith(".xlsx"):
        return _read_xlsx(content)

    text = content.decode("utf-8-sig")
    if name.endswith(".csv"):
        return [dict(row) for row in csv.DictReader(io.StringIO(text))][:MAX_ROWS]
    if name.endswith(".json") or name.endswith(".jsonl"):
        if name.endswith(".jsonl"):
            records = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            payload = json.loads(text)
            records = payload if isinstance(payload, list) else [payload]
        if not all(isinstance(row, dict) for row in records):
            raise ValueError("JSON analysis expects an object or a list of objects")
        return [dict(row) for row in records][:MAX_ROWS]
    raise ValueError("supported analysis formats are CSV, JSON, JSONL, and XLSX")


def analyze_payload(filename: str, content: bytes) -> dict[str, Any]:
    if not content:
        raise ValueError("analysis input cannot be empty")
    if len(content) > MAX_ANALYSIS_BYTES:
        raise ValueError("analysis input is too large")

    records = _parse_records(filename, content)
    summary = _summarize_records(records)
    duplicate_rows = _duplicate_rows(records)
    return {
        "filename": filename,
        "rows": len(records),
        "columns": sorted({str(key) for row in records for key in row}),
        "duplicate_rows": duplicate_rows,
        "missing_cells": sum(item["missing"] for item in summary.values()),
        "summary": summary,
        "sample": records[:MAX_SAMPLE_ROWS],
        "charts": _chart_specs(records),
    }
