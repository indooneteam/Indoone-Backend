from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfWriter
import pytest

from app.capabilities.document_extract import DOCX_MIME, MAX_PDF_PAGES, PDF_MIME, extract_document


def test_extract_docx_text() -> None:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Indoone document understanding")
    document.add_paragraph("Kannada ಕನ್ನಡ and Hindi हिंदी")
    document.save(buffer)

    result = extract_document("sample.docx", DOCX_MIME, buffer.getvalue())

    assert result.page_count is None
    assert result.paragraphs == 2
    assert "Indoone document understanding" in result.text
    assert "ಕನ್ನಡ" in result.text
    assert result.truncated is False
    assert result.tables == []


def test_extract_docx_table() -> None:
    buffer = io.BytesIO()
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Language"
    table.cell(0, 1).text = "Code"
    table.cell(1, 0).text = "Kannada"
    table.cell(1, 1).text = "kn"
    document.save(buffer)

    result = extract_document("table.docx", DOCX_MIME, buffer.getvalue())

    assert result.tables == [[["Language", "Code"], ["Kannada", "kn"]]]


def test_extract_pdf_text() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(buffer)

    result = extract_document("sample.pdf", PDF_MIME, buffer.getvalue())

    assert result.page_count == 1
    assert result.mime_type == PDF_MIME
    assert result.text == ""
    assert result.truncated is False
    assert result.tables == []


def test_reject_unsupported_document() -> None:
    with pytest.raises(ValueError, match="unsupported document type"):
        extract_document("sample.txt", "text/plain", b"hello")


def test_reject_path_traversal_filename() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(buffer)

    result = extract_document("../../sample.pdf", PDF_MIME, buffer.getvalue())
    assert result.filename == "sample.pdf"


def test_reject_windows_style_path_filename() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(buffer)

    result = extract_document(r"..\uploads\sample.pdf", PDF_MIME, buffer.getvalue())
    assert result.filename == "sample.pdf"


def test_reject_control_character_filename() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(buffer)

    with pytest.raises(ValueError, match="invalid document filename"):
        extract_document("sample\n.pdf", PDF_MIME, buffer.getvalue())


def test_reject_malformed_pdf() -> None:
    with pytest.raises(ValueError, match="invalid PDF document"):
        extract_document("broken.pdf", PDF_MIME, b"not a real pdf")


def test_reject_pdf_with_wrong_signature() -> None:
    with pytest.raises(ValueError, match="invalid PDF document"):
        extract_document("wrong.pdf", PDF_MIME, b"PK\x03\x04not a pdf")


def test_reject_malformed_docx() -> None:
    with pytest.raises(ValueError, match="invalid DOCX document"):
        extract_document("broken.docx", DOCX_MIME, b"not a zip document")


def test_reject_zip_that_is_not_docx() -> None:
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("payload.txt", "not a docx")

    with pytest.raises(ValueError, match="invalid DOCX document"):
        extract_document("wrong.docx", DOCX_MIME, buffer.getvalue())


def test_reject_docx_archive_that_expands_beyond_limit() -> None:
    buffer = io.BytesIO()
    payload = b"A" * 17_000_000
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", payload)

    with pytest.raises(ValueError, match="expands beyond safety limit"):
        extract_document("large.docx", DOCX_MIME, buffer.getvalue())


def test_reject_docx_with_too_many_table_cells(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "app.capabilities.document_extract.MAX_TOTAL_TABLE_CELLS",
        3,
    )
    buffer = io.BytesIO()
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    table.cell(1, 0).text = "c"
    table.cell(1, 1).text = "d"
    document.save(buffer)

    with pytest.raises(ValueError, match="too many table cells"):
        extract_document("large-table.docx", DOCX_MIME, buffer.getvalue())


def test_reject_pdf_that_exceeds_page_limit() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    for _ in range(MAX_PDF_PAGES + 1):
        writer.add_blank_page(width=300, height=300)
    writer.write(buffer)

    with pytest.raises(ValueError, match="PDF exceeds page limit"):
        extract_document("too-many-pages.pdf", PDF_MIME, buffer.getvalue())
